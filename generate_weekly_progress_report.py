"""Generate the 29 September 2026 MathRepair weekly progress DOCX."""

from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "MATHREPAIR_WEEKLY_PROGRESS_29_SEPTEMBER_2026.docx"
LEARNED_REPORT = ROOT / "learned_verifier_report.json"
BASELINE_REPORT = ROOT / "math500_full40_baseline_report.json"
SYSTEM_ABLATION_REPORT = ROOT / "math500_system_ablation_report.json"
VERIFIER_ABLATION_REPORT = ROOT / "verifier_ablation_report.json"
CATEGORY_REPORT = ROOT / "error_category_analysis.json"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EDF4FA"
GREEN = "548235"
PALE_GREEN = "E2F0D9"
AMBER = "BF8F00"
PALE_AMBER = "FFF2CC"
RED = "C00000"
PALE_RED = "FCE4D6"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"


def load_json(path: Path) -> dict[str, object]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"Expected a JSON object: {path}")
    return data


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_cell_text(cell, text: object, *, bold: bool = False, color: str | None = None, size: float = 8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, (cell, header) in enumerate(zip(header_row.cells, headers)):
        set_cell_text(cell, header, bold=True, color=WHITE, size=8.5)
        set_cell_shading(cell, NAVY)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, (cell, value) in enumerate(zip(cells, values)):
            set_cell_text(cell, value)
            if row_index % 2:
                set_cell_shading(cell, LIGHT_GRAY)
            if widths:
                cell.width = Inches(widths[column_index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 25, NAVY),
        ("Subtitle", 12, GRAY),
        ("Heading 1", 16, NAVY),
        ("Heading 2", 12, BLUE),
        ("Heading 3", 10.5, GREEN),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "MATHREPAIR  |  WEEKLY RESEARCH PROGRESS"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    add_page_number(section.footer.paragraphs[0])


def add_accent_rule(document: Document, color: str = BLUE) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    properties.append(borders)


def add_callout(document: Document, title: str, text: str, color: str, fill: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(color)
    detail = cell.add_paragraph(text)
    detail.paragraph_format.space_after = Pt(0)
    detail.runs[0].font.name = "Aptos"
    detail.runs[0].font.size = Pt(9.3)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def add_command_block(document: Document, commands: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    cell.text = ""
    for index, command in enumerate(commands):
        paragraph = cell.add_paragraph() if index else cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(3 if index < len(commands) - 1 else 0)
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def make_ablation_chart(report: dict[str, object], output: Path) -> None:
    variants = report["variants"]
    labels = [
        "Full",
        "No graph",
        "No typed\nerror",
        "No symbolic\ntool",
    ]
    initial = [100 * float(item["initial_answer_accuracy"]) for item in variants]
    final = [100 * float(item["final_answer_accuracy"]) for item in variants]
    positions = list(range(len(labels)))
    width = 0.34
    figure, axis = plt.subplots(figsize=(7.1, 3.0), dpi=180)
    axis.bar([value - width / 2 for value in positions], initial, width, label="Initial", color="#2F75B5")
    axis.bar([value + width / 2 for value in positions], final, width, label="After repair", color="#C00000")
    axis.set_ylabel("Answer accuracy (%)")
    axis.set_xticks(positions, labels)
    axis.set_ylim(0, max(12, max(initial + final) + 2))
    axis.grid(axis="y", color="#D9D9D9", linewidth=0.7)
    axis.set_axisbelow(True)
    axis.spines[["top", "right"]].set_visible(False)
    # Keep the legend above the plot so it cannot cover the final bar label.
    axis.legend(frameon=False, ncol=2, loc="lower center", bbox_to_anchor=(0.5, 1.01))
    for position, value in zip(positions, initial):
        axis.text(
            position - width / 2,
            value + 0.25,
            f"{value:.1f}%",
            ha="center",
            va="bottom",
            fontsize=8,
            color="#17365D",
            fontweight="bold",
        )
    for position, value in zip(positions, final):
        if value > 0:
            axis.text(
                position + width / 2,
                max(value * 0.5, 0.45),
                f"{value:.1f}%",
                ha="center",
                va="center",
                fontsize=8,
                color="white",
                fontweight="bold",
            )
        else:
            axis.text(
                position + width / 2,
                0.25,
                "0.0%",
                ha="center",
                va="bottom",
                fontsize=8,
                color="#C00000",
                fontweight="bold",
            )
    figure.tight_layout()
    figure.savefig(output, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def add_cover(document: Document) -> None:
    document.add_paragraph().paragraph_format.space_after = Pt(22)
    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = eyebrow.add_run("WEEKLY RESEARCH PROGRESS UPDATE")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    title = document.add_paragraph(style="Title")
    title.add_run("MathRepair")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Learned Typed Verification, MATH-500 Evaluation, and Matched-Budget Ablations"
    )
    add_accent_rule(document)

    metadata = document.add_table(rows=4, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.LEFT
    metadata.style = "Table Grid"
    for row, (label, value) in zip(
        metadata.rows,
        (
            ("Reporting date", "29 September 2026"),
            ("Research stage", "Supervisor follow-up experiments"),
            ("Benchmark", "MATH-500 hard subset, 40 level-4/5 problems"),
            ("Primary model", "qwen2.5:3b via local Ollama"),
        ),
    ):
        set_cell_text(row.cells[0], label, bold=True, color=NAVY, size=9)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[1], value, size=9)

    document.add_paragraph().paragraph_format.space_after = Pt(10)
    add_callout(
        document,
        "Purpose of this update",
        (
            "This document reports the three requested follow-up tasks: training the "
            "first learned typed verifier, performing an answer-scored evaluation on "
            "the full MATH-500 hard subset, and completing the no-graph, no-typed-error, "
            "and no-symbolic/tool system ablations."
        ),
        NAVY,
        PALE_BLUE,
    )
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(22)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Research artifacts and raw traces remain local pending supervisor review.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    note.add_run().add_break(WD_BREAK.PAGE)


def build_document() -> Document:
    learned = load_json(LEARNED_REPORT)
    baseline = load_json(BASELINE_REPORT)
    system = load_json(SYSTEM_ABLATION_REPORT)
    verifier_ablation = load_json(VERIFIER_ABLATION_REPORT)
    category_report = load_json(CATEGORY_REPORT)

    document = Document()
    document.core_properties.title = "MathRepair Weekly Research Progress Update"
    document.core_properties.subject = (
        "Learned typed verification, MATH-500 evaluation, and matched-budget ablations"
    )
    document.core_properties.author = "MathRepair Research Project"
    document.core_properties.keywords = (
        "MathRepair, MATH-500, learned verifier, local repair, ablation"
    )
    configure_document(document)
    add_cover(document)

    document.add_heading("Executive Summary", level=1)
    add_bullet(
        document,
        "The learned verifier achieved 100.0% location and error-type accuracy on an eight-example synthetic held-out test, while the current structural/symbolic baseline covered 50.0%.",
    )
    add_bullet(
        document,
        "The full 40-problem MATH-500 baseline obtained 10.0% answer accuracy (4/40), with 30 completed generations and 10 output failures retained in the denominator.",
    )
    add_bullet(
        document,
        "Every matched-budget system variant finished at 5.0% final accuracy. Full MathRepair corrected one wrong answer but regressed three initially correct answers.",
    )
    add_bullet(
        document,
        "The results do not support an accuracy advantage for the current local-repair configuration; they identify verifier generalization and repair acceptance as the immediate research problems.",
    )

    document.add_heading("Supervisor Deliverable Checklist", level=1)
    add_table(
        document,
        ["Requested deliverable", "Included in this report", "Primary artifact"],
        [
            ["Learned verifier results", "Yes: train/dev/test metrics and baseline comparison", "learned_verifier_report.json"],
            ["Real MATH-500 pilot accuracy", "Yes: all 40 problems, failures, calls and tokens", "math500_full40_baseline_report.json"],
            ["Completed key ablations", "Yes: no graph, no typed error, no symbolic/tool", "math500_system_ablation_report.json"],
            ["Error-type local-vs-global comparison", "Yes: category-wise success rates and observed winner", "error_category_analysis.json"],
            ["Short current conclusion", "Yes: current conclusion below", "This document, Scientific Conclusions"],
        ],
        widths=[2.15, 3.05, 1.75],
    )
    add_callout(
        document,
        "Scope followed",
        (
            "The main effort was spent on methodological contribution and experimental "
            "evidence. No new web-interface work was prioritized during this cycle."
        ),
        NAVY,
        PALE_BLUE,
    )

    add_callout(
        document,
        "Main scientific finding",
        (
            "Strong performance on controlled synthetic verifier data did not transfer "
            "to hard open-ended mathematical traces. The negative result is retained "
            "as evidence rather than hidden or optimized away."
        ),
        RED,
        PALE_RED,
    )

    document.add_heading("Requested Work and Completion Status", level=1)
    add_table(
        document,
        ["Focus area", "Status", "Evidence produced"],
        [
            [
                "Learned typed verifier",
                "Completed",
                "Deterministic training, development and test split; location/type metrics; rule/symbolic baseline comparison.",
            ],
            [
                "MATH-500 hard-subset evaluation",
                "Completed",
                "All 40 problems attempted; answer accuracy, output behavior, verifier detections, calls and token costs reported.",
            ],
            [
                "System-level ablations",
                "Completed",
                "No graph, no typed-error information, and no symbolic/tool support under a shared additional-token ceiling.",
            ],
        ],
        widths=[1.55, 1.05, 4.2],
    )

    document.add_heading("Learned Typed Verifier", level=1)
    document.add_paragraph(
        "The first learned verifier is an inspectable TF-IDF and logistic-regression baseline. "
        "One classifier localizes the first erroneous node and a second classifier predicts one "
        "of the eight frozen error types. Whole traces, rather than individual nodes, were kept "
        "disjoint across the split to reduce direct leakage."
    )
    split = learned["split"]
    document.add_paragraph(
        f"The controlled dataset contains {learned['dataset']['example_count']} examples: "
        f"{split['train']['count']} train, {split['dev']['count']} development, and "
        f"{split['test']['count']} test examples. The location threshold "
        f"({learned['location_threshold_selected_on_dev']:.2f}) was selected on the development split."
    )
    test = learned["test_metrics"]
    rule = learned["test_rule_baseline"]
    add_table(
        document,
        ["Method", "Coverage", "Location accuracy", "Type accuracy", "End-to-end"],
        [
            [
                "Learned typed verifier",
                "100.0%",
                f"{100 * test['error_location_accuracy']:.1f}%",
                f"{100 * test['error_type_accuracy']:.1f}%",
                f"{100 * test['end_to_end_accuracy']:.1f}%",
            ],
            [
                "Current structural/symbolic heuristics",
                f"{100 * rule['coverage']:.1f}%",
                f"{100 * rule['error_location_accuracy_with_abstentions_wrong']:.1f}%",
                f"{100 * rule['error_type_accuracy_with_abstentions_wrong']:.1f}%",
                "Not reported",
            ],
        ],
        widths=[2.25, 0.9, 1.2, 1.05, 1.0],
    )
    add_callout(
        document,
        "Interpretation boundary",
        (
            "The test contains only eight controlled synthetic examples, one per class. "
            "The 100.0% result validates the training pipeline but does not establish "
            "generalization to natural model errors."
        ),
        AMBER,
        PALE_AMBER,
    )

    document.add_heading("Full MATH-500 Hard-Subset Evaluation", level=1)
    document.add_paragraph(
        "The frozen 40-problem subset contains level-4 and level-5 problems from all seven "
        "MATH subjects. The baseline used qwen2.5:3b, seed 42, temperature 0, a 1,024-token "
        "output cap, and no repair. Reference answers were used only after generation for scoring."
    )
    add_table(
        document,
        ["Metric", "Result", "Measurement note"],
        [
            ["Problems attempted", baseline["problem_count"], "All remain in the accuracy denominator"],
            ["Completed generations", baseline["completed_count"], "Strict or normalized parse succeeded"],
            ["Output failures", baseline["failure_count"], "Not removed from accuracy"],
            ["Strict-contract outputs", baseline["strict_output_contract_count"], "Exact requested steps-array schema"],
            ["Normalized recoveries", baseline["normalized_recovery_count"], "Documented parser recovery"],
            ["Correct answers", baseline["initial_answer_correct_count"], "Exact/symbolic answer scoring"],
            ["Answer accuracy", f"{100 * baseline['initial_answer_accuracy']:.1f}%", "4 correct out of 40"],
            ["Model calls", baseline["total_model_calls"], "One baseline call per problem"],
            ["Prompt + completion tokens", f"{baseline['total_tokens']:,}", "Ollama-reported token counts"],
        ],
        widths=[2.0, 1.2, 3.8],
    )
    document.add_paragraph(
        "The low score reflects both incorrect mathematical answers and output-contract failures. "
        "Strict compliance was 1/40; 29 usable generations required normalized recovery. These "
        "two behaviors remain separate in the report so parser recovery is not presented as strict success."
    )

    document.add_heading("Matched-Budget System Ablations", level=1)
    add_callout(
        document,
        "Supervisor comment addressed",
        (
            "These ablations are important because they identify which part of "
            "MathRepair is actually useful. The current runs therefore keep the "
            "baseline generations, seeds, model, answer scorer, and additional-token "
            "ceiling fixed across variants. The observed differences are reported "
            "below, including null results and repair regressions."
        ),
        NAVY,
        PALE_BLUE,
    )
    document.add_paragraph(
        "Every arm reused the same frozen baseline generations. Full MathRepair established a "
        f"shared additional-token ceiling of {system['matched_additional_token_budget']:,}. "
        "Each detected case received at most one repair call; model, seed schedule, temperature, "
        "problem order, and answer scorer were held constant."
    )
    ablation_rows = []
    names = {
        "full_mathrepair": "Full MathRepair",
        "no_graph": "No graph",
        "no_typed_error": "No typed error",
        "no_symbolic_tool": "No symbolic/tool support",
    }
    for variant in system["variants"]:
        ablation_rows.append(
            [
                names[variant["variant_id"]],
                f"{100 * variant['initial_answer_accuracy']:.1f}%",
                f"{100 * variant['final_answer_accuracy']:.1f}%",
                variant["verifier_detection_count"],
                f"{variant['repair_success_count']}/{variant['repair_attempt_count']}",
                variant["repair_regression_count"],
                variant["additional_model_calls"],
                f"{variant['additional_tokens']:,}",
            ]
        )
    add_table(
        document,
        ["Variant", "Initial", "Final", "Detected", "Repair", "Regressions", "Calls", "Tokens"],
        ablation_rows,
        widths=[1.6, 0.65, 0.65, 0.65, 0.7, 0.75, 0.55, 0.8],
    )

    with TemporaryDirectory() as directory:
        chart = Path(directory) / "system_ablation_accuracy.png"
        make_ablation_chart(system, chart)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(chart), width=Inches(6.55))
    caption = document.add_paragraph(
        "Figure 1. Initial and post-repair answer accuracy under the matched system-ablation protocol."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8.5)
    caption.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

    document.add_heading("Ablation Interpretation", level=2)
    add_bullet(
        document,
        "Full MathRepair corrected one wrong answer in 20 repair attempts, but three initially correct answers became incorrect. Final accuracy fell from 10.0% to 5.0%.",
    )
    add_bullet(
        document,
        "Removing graph features increased detections from 20 to 30. Eleven calls were skipped because the corresponding full-arm per-problem allocation was zero; the arm corrected no answer and regressed two.",
    )
    add_bullet(
        document,
        "Removing typed-error information did not change final accuracy in this run. This null result is not evidence that typing is generally unnecessary because the model and verifier were already operating at low accuracy.",
    )
    add_bullet(
        document,
        "The structural/symbolic heuristic abstained on every completed open-ended trace. Therefore, the identical full and no-symbolic results demonstrate zero tool coverage on this path, not that symbolic support has no value.",
    )

    document.add_heading("Error-Category Analysis", level=1)
    document.add_paragraph(
        "This analysis addresses the central research question: when does local repair help, "
        "and when does global regeneration work better? The comparison unit is an accepted, "
        "answer-correct repair among detected-error runs in the earlier matched-budget algebra "
        "pilot. The adaptive local arm is the primary local-repair comparison; uniform local "
        "results are shown as a secondary control."
    )
    category_rows = []
    category_names = {
        "arithmetic_error": "Arithmetic error",
        "algebraic_transformation_error": "Algebraic transformation error",
        "missing_assumption": "Missing assumption",
        "dependency_error": "Dependency error",
        "incomplete_solution": "Incomplete reasoning",
        "sign_error": "Sign error",
        "logical_inference_error": "Logical inference error",
        "semantic_interpretation_error": "Semantic interpretation error",
    }
    for row in category_report["rows"]:
        cases = row["observed_error_runs"]
        category_rows.append(
            [
                category_names.get(row["error_category"], row["error_category"]),
                cases,
                f"{row['global_success_count']}/{cases}"
                if cases
                else "--",
                f"{row['uniform_local_success_count']}/{cases}"
                if cases
                else "--",
                f"{row['adaptive_local_success_count']}/{cases}"
                if cases
                else "--",
                row["best_observed_strategy"] or "--",
                row["evidence_status"],
            ]
        )
    add_table(
        document,
        ["Error category", "Cases", "Global success", "Uniform local", "Adaptive local", "Best observed", "Evidence"],
        category_rows,
        widths=[1.65, 0.5, 0.85, 0.8, 0.85, 1.0, 0.85],
    )
    add_callout(
        document,
        "Category-level finding",
        (
            "Global regeneration was stronger on the five observed algebraic-transformation "
            "cases (40.0% versus 20.0% for adaptive local repair). Adaptive local repair "
            "matched global regeneration on the single observed sign-error case (100.0%), "
            "while neither strategy succeeded on the single arithmetic-error case. "
            "Missing-assumption, dependency, incomplete-reasoning, logical-inference, and "
            "semantic-interpretation errors were not observed in this pilot."
        ),
        GREEN,
        PALE_GREEN,
    )
    document.add_paragraph(
        "This pattern supports the motivation for typed repair policies, but it is not yet a "
        "general claim: seven detected errors are too few, and five belong to one category. "
        "The next category-focused dataset should deliberately include valid traces and the "
        "currently unobserved error types."
    )

    document.add_heading("Verifier-Stage Ablation Check", level=2)
    verifier_rows = []
    for variant in verifier_ablation["variants"]:
        metrics = variant["metrics"]
        verifier_rows.append(
            [
                names.get(variant["variant_id"], variant["variant_id"]),
                f"{100 * metrics['error_location_accuracy']:.1f}%",
                "--" if metrics["error_type_accuracy"] is None else f"{100 * metrics['error_type_accuracy']:.1f}%",
                "--" if metrics["end_to_end_accuracy"] is None else f"{100 * metrics['end_to_end_accuracy']:.1f}%",
            ]
        )
    add_table(
        document,
        ["Variant", "Location", "Type", "End-to-end"],
        verifier_rows,
        widths=[2.7, 1.25, 1.25, 1.25],
    )
    document.add_paragraph(
        "These eight-example synthetic ablations are supporting diagnostics only. They are not "
        "substitutes for final-answer or repair outcomes on MATH-500."
    )

    document.add_heading("Scientific Conclusions", level=1)
    add_callout(
        document,
        "Current conclusion",
        (
            "The present evidence does not show that verifier-guided local repair "
            "outperforms a non-repair baseline on hard open-ended mathematics. "
            "The principal observed failure is unsafe repair acceptance after a "
            "synthetic-data verifier predicts errors outside its training distribution."
        ),
        RED,
        PALE_RED,
    )
    document.add_paragraph(
        "What currently works: the graph-integrated pipeline, reproducible typed-verifier "
        "training/evaluation, matched-budget accounting, and category-level comparison are "
        "operational. What still fails: the synthetic learned verifier transfers poorly to "
        "hard open-ended MATH traces, output-contract compliance is unstable, symbolic tools "
        "have zero coverage on that path, and local repair can regress an initially correct "
        "answer. The current methodological contribution is therefore the controlled framework "
        "for identifying these conditions, not a claim that local repair is universally superior."
    )
    add_bullet(
        document,
        "The learned-verifier training code and evaluation path are operational, but the current supervision set is too small and templated for broad claims.",
    )
    add_bullet(
        document,
        "MATH-500 exposes output-contract instability: only one of 40 baseline outputs strictly followed the requested schema.",
    )
    add_bullet(
        document,
        "A detector should learn a valid/no-repair outcome, and generated repairs need an independent acceptance check before replacing an initially correct answer.",
    )
    add_bullet(
        document,
        "The earlier matched-budget result in which global regeneration reached 88.9% and local repair reached 86.1% remains scientifically relevant; the new hard-subset result strengthens the need for failure analysis rather than result tuning.",
    )

    document.add_heading("Limitations and Pending Direction", level=1)
    add_table(
        document,
        ["Limitation", "Consequence", "Candidate response after supervisor review"],
        [
            [
                "Synthetic verifier data",
                "Perfect held-out template performance does not transfer to natural MATH traces.",
                "Create reviewed natural-error and valid-trace labels.",
            ],
            [
                "No learned valid/no-repair class",
                "The verifier can trigger repair on correct traces.",
                "Train calibrated detection with abstention and valid examples.",
            ],
            [
                "No independent repair acceptance for open text",
                "Correct answers can regress after local generation.",
                "Add a separate acceptance verifier or answer-consistency gate.",
            ],
            [
                "Symbolic coverage is zero on open-ended traces",
                "The no-symbolic ablation is non-informative.",
                "Introduce domain-aware tool calls or formalizable substeps.",
            ],
            [
                "Small local model and unstable output contract",
                "Ten baseline failures and low overall accuracy.",
                "Confirm the next model/output protocol before further large runs.",
            ],
        ],
        widths=[1.7, 2.45, 2.65],
    )
    document.add_paragraph(
        "No new methodological direction is started in this update. The candidate responses above "
        "are held for supervisor prioritization."
    )

    document.add_heading("Reproducibility and Evidence", level=1)
    document.add_paragraph(
        "The commands below regenerate the learned-verifier report, run or resume the full baseline, "
        "run the matched system ablations, and execute the automated checks. Raw JSONL outputs are "
        "preserved locally and excluded from Git because they are large regenerable artifacts."
    )
    add_command_block(
        document,
        [
            "python learned_verifier.py",
            "python math500_evaluation.py --model qwen2.5:3b --repair-attempts 0 --timeout-seconds 120 --max-output-tokens 1024 --report math500_full40_baseline_report.json --traces math500_full40_baseline_traces.jsonl",
            "python math500_evaluation.py --model qwen2.5:3b --repair-attempts 0 --timeout-seconds 120 --max-output-tokens 1024 --report math500_full40_baseline_report.json --traces math500_full40_baseline_traces.jsonl --resume",
            "python math500_system_ablation.py --model qwen2.5:3b --timeout-seconds 60 --max-output-tokens 192 --resume",
            "python show_experiments.py",
            "python -m unittest discover",
        ],
    )
    add_table(
        document,
        ["Artifact", "Purpose"],
        [
            ["learned_verifier_report.json", "Train/dev/test metrics and baseline comparison"],
            ["math500_full40_baseline_report.json", "Full hard-subset accuracy and output behavior"],
            ["math500_system_ablation_report.json", "Matched-budget system-level ablation results"],
            ["verifier_ablation_report.json", "Controlled verifier-stage ablations"],
            ["SUPERVISOR_NEXT_WEEK_RESULTS.md", "Concise machine-readable-text summary"],
        ],
        widths=[2.7, 4.1],
    )
    document.add_paragraph(
        "Automated verification at the time of this report: 144/144 tests passing."
    ).runs[0].bold = True

    return document


def main() -> None:
    document = build_document()
    document.save(OUTPUT)
    print(f"Weekly progress report: {OUTPUT}")
    print(f"Paragraphs: {len(document.paragraphs)}; tables: {len(document.tables)}")


if __name__ == "__main__":
    main()
