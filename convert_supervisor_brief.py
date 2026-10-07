from pathlib import Path

from docx import Document
from docx.shared import Pt


source = Path(__file__).with_name("SUPERVISOR_FEEDBACK_BRIEF.md")
target = Path(__file__).with_name("SUPERVISOR_FEEDBACK_BRIEF.docx")
document = Document()

for line in source.read_text(encoding="utf-8").splitlines():
    if not line.strip():
        document.add_paragraph()
    elif line.startswith("# "):
        document.add_heading(line[2:].strip(), level=1)
    elif line.startswith("## "):
        document.add_heading(line[3:].strip(), level=2)
    elif line.startswith("- "):
        document.add_paragraph(line[2:].strip(), style="List Bullet")
    elif line[:2].isdigit() or (line[:1].isdigit() and line[1:2] == "."):
        document.add_paragraph(line.split(". ", 1)[-1], style="List Number")
    elif line.startswith("|"):
        continue
    else:
        document.add_paragraph(line)

for paragraph in document.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(11)

document.save(target)
print(target)
