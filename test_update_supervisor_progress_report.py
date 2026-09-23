import unittest

from docx import Document

from update_supervisor_progress_report import (
    UPDATE_HEADING,
    remove_existing_update,
    replace_next_steps_section,
)


class SupervisorProgressReportTests(unittest.TestCase):
    def test_removes_existing_update_and_every_element_after_it(self):
        document = Document()
        document.add_paragraph("Original report")
        document.add_heading(UPDATE_HEADING, level=1)
        document.add_paragraph("Old update")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Old table"
        document.add_heading(UPDATE_HEADING, level=1)

        remove_existing_update(document)

        self.assertEqual(
            [paragraph.text for paragraph in document.paragraphs],
            ["Original report"],
        )
        self.assertEqual(len(document.tables), 0)

        document.add_heading("10. Recommended Next Steps", level=1)
        document.add_paragraph("Old recommendation")
        document.add_heading("11. Demonstration Commands", level=1)
        replace_next_steps_section(document)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("10.1 Completed Supervisor Tasks", text)
        self.assertIn("10.2 Explicitly Deferred Work", text)
        self.assertIn("10.3 Proposed Follow-Up Experiments", text)
        self.assertNotIn("Old recommendation", text)


if __name__ == "__main__":
    unittest.main()
