"""Update the supervisor DOCX with the completed methodology tasks."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


REPORT_PATH = Path(__file__).with_name("MATHREPAIR_SUPERVISOR_PROGRESS_REPORT.docx")
UPDATE_HEADING = "12. Supervisor Tasks 1-7 Completion Update"

TASK_ROWS = [
    (
        "1",
        "Completed",
        "Real model traces now use the reasoning graph internally. Each node stores state, parents, verification, error impact, repair, and final status.",
        "model_pipeline.py; reasoning_graph.py",
    ),
    (
        "2",
        "Completed within requested scope",
        "A deterministic 40-problem MATH-500 subset has 17 level-4 and 23 level-5 problems across seven subjects. The smoke test processed 40/40 in text_unverified mode.",
        "math500_pilot_40_manifest.json",
    ),
    (
        "3",
        "Completed",
        "Eight typed errors are frozen with definitions, positive and negative examples, detection contracts, and repair actions.",
        "ERROR_TAXONOMY.md",
    ),
    (
        "4",
        "Completed within requested scope",
        "The data generator produced 48 typed-verifier examples, six per error type, with error location, corrupted/correct steps, and preferred action.",
        "typed_verifier_pilot.json",
    ),
    (
        "5",
        "Completed",
        "Under a shared approximately 5,800-token ceiling, global regeneration reached 88.9% answer accuracy; uniform and adaptive local repair each reached 86.1%. Budget violations: zero.",
        "matched_budget_report.md",
    ),
    (
        "6",
        "Completed within requested scope",
        "The requested six-row structure is frozen. As explicitly allowed, two rows are measured, one is a rule-based proxy reference, and three future cells remain unfilled.",
        "ablation_table.md",
    ),
    (
        "7",
        "Completed",
        "The archive preserves exact model digests, prompt/parser versions, 432 raw-output records, generation failures, trace hashes, and normalized accuracy for four configurations.",
        "output_contract_evidence.md",
    ),
]


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8)


def remove_existing_update(document: Document) -> None:
    body = document._element.body
    found = False
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        text = Paragraph(child, document).text if child.tag == qn("w:p") else ""
        if text.strip() == UPDATE_HEADING:
            found = True
        if found:
            body.remove(child)


def replace_status_text(document: Document) -> None:
    replacements = {
        "Research progress review | September 2026": (
            "Research progress review | Updated 23 September 2026\n"
            "Repository: https://github.com/mwanajh/MathRepair_Project"
        ),
        "This work implements and evaluates a MathRepair prototype": (
            "MathRepair now connects model-generated traces directly to its reasoning "
            "graph and includes a 40-problem MATH-500 pilot, a frozen eight-class "
            "error taxonomy, a 48-example typed-verifier dataset, a fresh three-arm "
            "matched-budget pilot, a frozen ablation matrix, and a versioned archive "
            "of 432 raw outputs. The earlier expanded algebra experiment found "
            "100.0% local-repair accuracy versus 92.6% global regeneration. A newer "
            "and stricter total-budget pilot found 88.9% for global regeneration and "
            "86.1% for both local arms. These protocols answer different questions, "
            "so the report does not claim universal local-repair superiority."
        ),
        "Defensible claim:": (
            "Defensible claim: The proposal-level experimental infrastructure is now "
            "implemented and produces reproducible, mixed results. Local repair led "
            "in the earlier expanded pilot, while global regeneration led by 2.8 "
            "percentage points in the fresh total-token-matched pilot."
        ),
        "Executable symbolic verifier and repair pipeline with": (
            "Executable graph-based verifier and repair pipeline with 129 passing tests."
        ),
        "Executable graph-based verifier and repair pipeline with": (
            "Executable graph-based verifier and repair pipeline with 129 passing tests."
        ),
        "The complete proposal experiment is not yet finished.": (
            "The complete proposal experiment is not yet finished. The project now "
            "includes a 40-problem MATH-500 level-4/5 processing pilot, a frozen "
            "six-row ablation design, and graph-integrated model traces. However, the "
            "MATH-500 run is currently a text-processing smoke test rather than an "
            "accuracy evaluation; the learned calibrated verifier is not trained; "
            "three ablation cells remain unmeasured; and the fresh matched-budget "
            "result is based on only 36 runs and seven detected errors."
        ),
    }
    for paragraph in document.paragraphs:
        original = paragraph.text.strip()
        for prefix, replacement in replacements.items():
            if original.startswith(prefix):
                paragraph.text = replacement
                break


def insert_table_before(document: Document, paragraph: Paragraph, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.rows[0].cells[0], "Field", bold=True)
    set_cell_text(table.rows[0].cells[1], "Protocol", bold=True)
    for field, protocol in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], field)
        set_cell_text(cells[1], protocol)
    paragraph._element.addprevious(table._element)


def replace_experimental_design_section(document: Document) -> None:
    """Document the earlier and current protocols without mixing estimands."""
    body = document._element.body
    children = list(body)
    start = next(
        index
        for index, child in enumerate(children)
        if child.tag == qn("w:p")
        and Paragraph(child, document).text.strip().startswith("5. Experimental Design")
    )
    end = next(
        index
        for index, child in enumerate(children[start + 1 :], start + 1)
        if child.tag == qn("w:p")
        and Paragraph(child, document).text.strip().startswith("6. ")
    )
    design_heading = Paragraph(children[start], document)
    results_heading = Paragraph(children[end], document)
    design_heading.text = "5. Experimental Design and Protocol Separation"
    design_heading.style = "Heading 1"
    results_heading.text = "6. Earlier Expanded Evaluation Results"
    results_heading.style = "Heading 1"

    for child in children[start + 1 : end]:
        body.remove(child)

    paragraph = results_heading.insert_paragraph_before(
        "The report contains two complementary experimental protocols. They use "
        "different budget definitions and must not be pooled into one performance claim."
    )
    paragraph.style = "Normal"

    paragraph = results_heading.insert_paragraph_before(
        "5.1 Earlier Expanded Repeated-Seed Evaluation"
    )
    paragraph.style = "Heading 2"
    insert_table_before(
        document,
        results_heading,
        [
            ("Purpose", "Earlier evidence for local versus global repair on symbolic algebra."),
            ("Model", "qwen2-math:1.5b through the local Ollama API."),
            ("Dataset", "36 algebra problems: 18 earlier held-out and 18 new."),
            ("Trials", "Three seeds and 36 runs per seed; 108 completed runs."),
            ("Strategies", "No repair, verified global regeneration, and verified local repair."),
            ("Budget", "Matched per detected-error case under the earlier protocol."),
            ("Metrics", "Answer accuracy, valid-trace rate, calls, tokens, and paired confidence intervals."),
        ],
    )

    paragraph = results_heading.insert_paragraph_before(
        "5.2 Supervisor Task 1-7 Methodology"
    )
    paragraph.style = "Heading 2"
    insert_table_before(
        document,
        results_heading,
        [
            ("Internal representation", "Model traces are verified and repaired as reasoning graphs with node state, parents, error impact, repair, and final status."),
            ("Hard pilot", "A deterministic 40-problem MATH-500 subset: 17 level-4 and 23 level-5 problems across seven subjects; processing check only."),
            ("Typed errors", "Eight frozen error classes with definitions, examples, detection contracts, and allowed repair actions."),
            ("Verifier supervision", "48 controlled corrupted/correct trace pairs, balanced at six examples per error type."),
            ("Fresh matched budget", "Thirty-six initial runs and seven detected-error cases; global, uniform-local, and adaptive-local arms share an approximately 5,800 additional-token ceiling."),
            ("Ablations", "Six frozen variants with explicit measured, proxy, planned, and learned-verifier-dependent evidence states."),
            ("Output contracts", "Four saved configurations preserve model digests, prompt/parser versions, 432 raw outputs, failures, and strict/normalized accuracy."),
            ("Primary fresh metric", "Answer accuracy over all scheduled runs; secondary metrics include valid traces, repair success, calls, tokens, and budget violations."),
        ],
    )

    paragraph = results_heading.insert_paragraph_before(
        "5.3 Interpretation Boundary"
    )
    paragraph.style = "Heading 2"
    results_heading.insert_paragraph_before(
        "The earlier expanded experiment reported 100.0% local repair versus "
        "92.6% global regeneration. The stricter fresh total-token-matched pilot "
        "reported 86.1% adaptive local repair versus 88.9% global regeneration. "
        "Because the protocols answer different questions, neither result replaces "
        "the other and no universal superiority claim is made."
    )


def replace_next_steps_section(document: Document) -> None:
    """Separate completed tasks from deferred and newly proposed work."""
    paragraphs = document.paragraphs
    start = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip()
        in {
            "10. Recommended Next Steps",
            "10. Work Status After Supervisor Tasks",
        }
    )
    end = next(
        index
        for index, paragraph in enumerate(paragraphs[start + 1 :], start + 1)
        if paragraph.text.strip().startswith("11. Demonstration Commands")
    )
    heading = paragraphs[start]
    next_section = paragraphs[end]
    heading.text = "10. Work Status After Supervisor Tasks"
    heading.style = "Heading 1"

    for paragraph in paragraphs[start + 1 : end]:
        paragraph._element.getparent().remove(paragraph._element)

    content = [
        (
            "10.1 Completed Supervisor Tasks",
            "Heading 2",
        ),
        (
            "Tasks 1-7 are completed within the scope requested for this week. "
            "The evidence and results are summarized in Section 12.",
            None,
        ),
        (
            "10.2 Explicitly Deferred Work",
            "Heading 2",
        ),
        (
            "Train and calibrate the learned typed verifier. Task 4 explicitly "
            "required the data-generation pipeline this week, not completion of "
            "large-model training.",
            "List Bullet",
        ),
        (
            "Run a full real-model MATH-500 accuracy evaluation. Task 2 explicitly "
            "limited this week to creating and processing a difficult 30-50 problem pilot.",
            "List Bullet",
        ),
        (
            "Fill the no-graph, no-typed-error, and no-symbolic-tool ablation cells. "
            "Task 6 explicitly requested the table structure even if cells were unfinished.",
            "List Bullet",
        ),
        (
            "These items are deferred by the stated scope; they are not incomplete "
            "deliverables from Tasks 1-7.",
            None,
        ),
        (
            "10.3 Proposed Follow-Up Experiments",
            "Heading 2",
        ),
        (
            "Repeat the three-arm matched-budget comparison on the hard benchmark "
            "pilot and report paired uncertainty.",
            "List Bullet",
        ),
        (
            "Add natural model errors to the controlled verifier data before final training.",
            "List Bullet",
        ),
        (
            "Keep prompt/parser sensitivity as supporting evidence rather than the primary contribution.",
            "List Bullet",
        ),
    ]
    for text, style in content:
        paragraph = next_section.insert_paragraph_before(text)
        if style:
            paragraph.style = style


def add_task_update(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    document.add_heading(UPDATE_HEADING, level=1)
    document.add_paragraph(
        "This section records the seven methodology tasks requested after the "
        "earlier progress review. It supersedes older status statements where "
        "the graph, benchmark pilot, and ablation plan were still pending."
    )

    task_table = document.add_table(rows=1, cols=4)
    task_table.style = "Table Grid"
    headers = ["Task", "Status", "Evidence / current result", "Artifact"]
    for cell, header in zip(task_table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
    for task, status, evidence, artifact in TASK_ROWS:
        cells = task_table.add_row().cells
        for cell, value in zip(cells, [task, status, evidence, artifact]):
            set_cell_text(cell, value)

    document.add_heading("12.1 Fresh Matched-Budget Result", level=2)
    document.add_paragraph(
        "The fresh Task 5 run used 36 model calls as the initial sample, found "
        "seven error cases, and applied an approximately 5,800 additional-token "
        "ceiling to each repair strategy."
    )
    budget_table = document.add_table(rows=1, cols=6)
    budget_table.style = "Table Grid"
    headers = ["Strategy", "Answer", "Valid trace", "Repair", "Avg calls", "Avg tokens"]
    rows = [
        ("No repair", "86.1%", "80.6%", "0.0%", "1.00", "624.1"),
        ("Global regeneration", "88.9%", "88.9%", "42.9%", "1.25", "757.4"),
        ("Uniform local", "86.1%", "83.3%", "14.3%", "1.25", "772.1"),
        ("Adaptive local", "86.1%", "86.1%", "28.6%", "1.17", "721.5"),
    ]
    for cell, header in zip(budget_table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
    for row in rows:
        cells = budget_table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    document.add_paragraph(
        "Interpretation: adaptive local repair improved trace validity over "
        "uniform allocation and used fewer actual tokens, but it did not improve "
        "answer accuracy. Global regeneration was 2.8 percentage points more "
        "accurate in this small pilot."
    )

    document.add_heading("12.2 Ablation and Learned-Verifier Status", level=2)
    document.add_paragraph(
        "The ablation matrix now contains Full MathRepair, no graph, no typed "
        "error, no adaptive compute, global regeneration instead of local repair, "
        "and no symbolic tool. Only no adaptive compute and global regeneration "
        "are currently marked measured. Full MathRepair is a rule-based proxy; "
        "the remaining cells stay explicit rather than receiving estimated values."
    )

    document.add_heading("12.3 Output-Contract Sensitivity", level=2)
    contract_table = document.add_table(rows=1, cols=4)
    contract_table.style = "Table Grid"
    headers = ["Configuration", "Failures", "Strict", "Normalized"]
    rows = [
        ("qwen2-math:1.5b / default", "0/108", "90.7%", "90.7%"),
        ("qwen2.5:3b / default", "24/108", "47.2%", "59.3%"),
        ("qwen2-math:7b / default", "91/108", "13.0%", "41.7%"),
        ("qwen2-math:7b / JSON contract", "10/108", "75.0%", "79.6%"),
    ]
    for cell, header in zip(contract_table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
    for row in rows:
        cells = contract_table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    document.add_paragraph(
        "Normalized accuracy is an offline parser-only sensitivity analysis and "
        "does not replace the strict end-to-end result. Exact model digests, "
        "prompt/parser versions, trace hashes, and 432 raw-output records are "
        "preserved in output_contract_evidence.json and its raw archive."
    )

    document.add_heading("12.4 Reproduction", level=2)
    for command in [
        "python show_experiments.py",
        "python generate_verifier_dataset.py --count-per-type 6 --seed 42",
        "python matched_budget_experiment.py --reuse-global-results --reuse-local-results",
        "python ablation_table.py",
        "python output_contract_evidence.py",
        "python -m unittest discover -q",
    ]:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    document.add_paragraph(
        "Current automated verification: 129/129 tests passing. The immediate "
        "research priority is the learned verifier and completion of the pending "
        "controlled ablations, not additional web-interface work."
    )


def main() -> None:
    document = Document(REPORT_PATH)
    remove_existing_update(document)
    replace_status_text(document)
    replace_experimental_design_section(document)
    replace_next_steps_section(document)
    add_task_update(document)

    section = document.sections[-1]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    document.save(REPORT_PATH)
    print(f"Updated: {REPORT_PATH}")


if __name__ == "__main__":
    main()
